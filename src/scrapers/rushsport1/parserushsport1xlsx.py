import io
from datetime import date, datetime, time

import pandas as pd
from openpyxl import load_workbook

DATE_COLUMNS = {"Program Start Date"}
TIME_COLUMNS = {"Program Start Time", "Program End Time"}


def parseRushSport1Xlsx(file_content: bytes) -> pd.DataFrame:
    workbook = load_workbook(io.BytesIO(file_content), data_only=True)
    worksheet = workbook.active

    rows = list(worksheet.iter_rows())
    if not rows:
        return pd.DataFrame()

    headers = [cell.value for cell in rows[0]]
    records = []

    for row in rows[1:]:
        if _row_is_empty(row):
            continue

        records.append(_build_record(headers, row))

    return pd.DataFrame(records)


def _build_record(headers: list, row) -> dict:
    record = {}

    for header, cell in zip(headers, row):
        if header is None:
            continue

        record[header] = _normalize_cell_value(header, cell)

    return record


def _normalize_cell_value(header: str, cell):
    value = cell.value
    if value is None:
        return None

    if header in DATE_COLUMNS:
        return _parse_date_value(value)

    if header in TIME_COLUMNS:
        return _parse_time_value(value)

    return value


def _parse_date_value(value):
    normalized_date = _coerce_date(value)
    if normalized_date is not None:
        return normalized_date

    return str(value).strip()


def _coerce_date(value) -> date | None:
    if isinstance(value, datetime):
        value = value.date()

    if not isinstance(value, date):
        return None

    return value


def _parse_time_value(value):
    normalized_time = _coerce_time(value)
    if normalized_time is not None:
        return normalized_time.replace(second=0, microsecond=0)

    return str(value).strip()


def _coerce_time(value) -> time | None:
    if isinstance(value, datetime):
        return value.time()

    if isinstance(value, time):
        return value

    return None


def _row_is_empty(row) -> bool:
    return all(cell.value is None for cell in row)
