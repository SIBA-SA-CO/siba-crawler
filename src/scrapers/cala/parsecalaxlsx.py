import io
from datetime import datetime, time

import pandas as pd
from openpyxl import load_workbook

TIME_COLUMNS = {"Start Time", "Duration"}


def parseCalaXlsx(file_content: bytes) -> pd.DataFrame:
    """Lee un Excel CALA desde bytes y devuelve sus filas como `DataFrame`.

    Los archivos del proveedor pueden contener varias hojas y no siempre ubican
    la cabecera en la primera fila. El parser inspecciona cada hoja, detecta la
    fila de encabezados por columnas obligatorias, normaliza valores de hora y
    aplica una correccion especifica para desfases AM/PM detectados en origen.
    """
    workbook = load_workbook(io.BytesIO(file_content), data_only=True)
    records = []

    for worksheet in workbook.worksheets:
        rows = list(worksheet.iter_rows())
        if not rows:
            continue

        header_row_index, headers = _find_header_row(rows)
        if header_row_index is None:
            continue

        for row in rows[header_row_index + 1:]:
            if _row_is_empty(row):
                continue

            record = _build_record(headers, row)
            if record:
                records.append(record)

    _correct_start_times(records)
    return pd.DataFrame(records)


def _find_header_row(rows) -> tuple[int | None, list]:
    """Localiza la fila de encabezados por las columnas minimas requeridas."""
    for index, row in enumerate(rows):
        headers = [cell.value for cell in row]
        normalized_headers = {str(header).strip() for header in headers if header is not None}

        if {"Date", "Start Time", "Title"}.issubset(normalized_headers):
            return index, headers

    return None, []


def _build_record(headers: list, row) -> dict:
    """Construye un diccionario columna-valor para una fila del Excel."""
    record = {}

    for header, cell in zip(headers, row):
        if header is None:
            continue

        normalized_header = str(header).strip()
        record[normalized_header] = _normalize_cell_value(normalized_header, cell.value)

    return record


def _normalize_cell_value(header: str, value):
    """Normaliza solo los tipos que necesitan tratamiento especial."""
    if value is None:
        return None

    if header in TIME_COLUMNS:
        return _format_time_value(value)

    return value


def _format_time_value(value):
    """Elimina segundos y microsegundos de horas leidas por openpyxl."""
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)

    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)

    return value


def _correct_start_times(records: list[dict]) -> None:
    """Corrige in situ horas de inicio que vienen desplazadas 12 horas.

    Algunos registros de CALA llegan con una hora visualmente valida, pero 12
    horas adelante respecto a la hora esperada segun el evento anterior y su
    duracion. Cuando se detecta exactamente ese desfase dentro de la misma fecha,
    se resta 12 horas al inicio del registro actual.
    """
    previousDate = None
    previousStart = None
    previousDuration = None

    for record in records:
        currentDate = record.get("Date")
        currentStart = record.get("Start Time")
        currentDuration = record.get("Duration")

        if currentDate != previousDate:
            previousDate = currentDate
            previousStart = currentStart if isinstance(currentStart, time) else None
            previousDuration = currentDuration if isinstance(currentDuration, time) else None
            continue

        if (
            isinstance(currentStart, time)
            and isinstance(previousStart, time)
            and isinstance(previousDuration, time)
        ):
            expectedStart = _add_duration(previousStart, previousDuration)
            if _seconds_since_midnight(currentStart) - _seconds_since_midnight(expectedStart) == 12 * 60 * 60:
                currentStart = _subtract_hours(currentStart, 12)
                record["Start Time"] = currentStart

        previousStart = currentStart if isinstance(currentStart, time) else None
        previousDuration = currentDuration if isinstance(currentDuration, time) else None


def _add_duration(start: time, duration: time) -> time:
    """Suma una duracion expresada como `time` a una hora de inicio."""
    total_seconds = (
        _seconds_since_midnight(start)
        + duration.hour * 3600
        + duration.minute * 60
        + duration.second
    ) % (24 * 60 * 60)
    return _time_from_seconds(total_seconds)


def _subtract_hours(value: time, hours: int) -> time:
    """Resta horas manteniendo el resultado dentro del rango de un dia."""
    total_seconds = (_seconds_since_midnight(value) - hours * 3600) % (24 * 60 * 60)
    return _time_from_seconds(total_seconds)


def _seconds_since_midnight(value: time) -> int:
    """Convierte una hora a segundos desde medianoche."""
    return value.hour * 3600 + value.minute * 60 + value.second


def _time_from_seconds(total_seconds: int) -> time:
    """Convierte segundos desde medianoche a `time` con wrap de 24 horas."""
    total_seconds = total_seconds % (24 * 60 * 60)
    hours = total_seconds // 3600
    minutes = (total_seconds % 3600) // 60
    seconds = total_seconds % 60
    return time(hour=hours, minute=minutes, second=seconds)


def _row_is_empty(row) -> bool:
    """Indica si una fila no contiene ningun valor util."""
    return all(cell.value is None for cell in row)
