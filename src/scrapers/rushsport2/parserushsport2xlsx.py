import io
from datetime import date, datetime, time

import pandas as pd
from openpyxl import load_workbook

DATE_COLUMNS = {"Program Start Date"}
TIME_COLUMNS = {"Program Start Time", "Program End Time"}


def parseRushSport2Xlsx(file_content: bytes) -> pd.DataFrame:
    workbook = load_workbook(io.BytesIO(file_content), data_only=True)
    worksheet = workbook.active

    rows = list(worksheet.iter_rows())
    if not rows:
        return pd.DataFrame()

    headers = [cell.value for cell in rows[0]]
    records = []

    for row in rows[1:]:
        if _row_is_empty(row):
            continue

        records.append(_build_record(headers, row))

    return pd.DataFrame(records)


def _build_record(headers: list, row) -> dict:
    record = {}

    for header, cell in zip(headers, row):
        if header is None:
            continue

        record[header] = _normalize_cell_value(header, cell.value)

    return record


def _normalize_cell_value(header: str, value):
    if value is None:
        return None

    if header in DATE_COLUMNS:
        normalizedDate = _coerce_date(value)
        if normalizedDate is not None:
            return normalizedDate

    if header in TIME_COLUMNS:
        normalizedTime = _coerce_time(value)
        if normalizedTime is not None:
            return normalizedTime.replace(second=0, microsecond=0)

    return str(value).strip() if isinstance(value, str) else value


def _coerce_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    return None


def _coerce_time(value) -> time | None:
    if isinstance(value, datetime):
        return value.time()

    if isinstance(value, time):
        return value

    return None


def _row_is_empty(row) -> bool:
    return all(cell.value is None for cell in row)

